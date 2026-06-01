from __future__ import annotations

import sys
import os

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import asyncio
import json
import logging
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, AsyncGenerator

from fastapi import FastAPI, HTTPException, Query, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse, Response, JSONResponse
from pydantic import BaseModel

from config import config

logger = logging.getLogger(__name__)


# ── Audit log helper ───────────────────────────────────────────────────────────

def append_audit(
    session: dict,
    phase: int,
    event_type: str,
    title: str,
    detail: str,
    data: dict | None = None,
    status: str = "info",
) -> None:
    try:
        session["audit_log"].append({
            "id": str(uuid.uuid4()),
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "phase": phase,
            "event_type": event_type,
            "title": title,
            "detail": detail,
            "data": data,
            "status": status,
        })
    except Exception:
        pass


# ── Session store ──────────────────────────────────────────────────────────────

sessions: dict[str, dict] = {}


def new_session() -> dict:
    return {
        "db": {
            "type": "postgresql",
            "host": "localhost",
            "port": 5432,
            "database": "",
            "username": "",
            "password": "",
        },
        "superset": {
            "url": config.SUPERSET_URL,
            "username": config.SUPERSET_USERNAME,
            "password": config.SUPERSET_PASSWORD,
            "session_cookie": "",
            "csrf_token": "",
        },
        "llm_model": config.LLM_MODEL,
        "db_connector": None,
        "phase1": {
            "schema_map": None,
            "confirmed": False,
            "business_prompt": "",
            "cache_key": None,
            "excluded_tables": {},
        },
        "phase2": {
            "query_plan": None,
            "qa_report": None,
            "confirmed": False,
            "edited_sql": "",
        },
        "phase3": {
            "dataset_name": "",
            "dashboard_title": "",
            "requirements": "",
            "dataset_info": None,
            "parsed_reqs": None,
            "dashboard_plan": None,
            "plan_ready": False,
            "dashboard_id": None,
            "dashboard_url": None,
            "chart_ids": [],
            "dq_report": None,
        },
        "audit_log": [],
        "csv_session": None,
    }


def get_session(sid: str) -> dict:
    if sid not in sessions:
        raise HTTPException(status_code=404, detail=f"Session '{sid}' not found")
    return sessions[sid]


def serialize_session(sess: dict) -> dict:
    """Return a JSON-serializable snapshot (no live objects)."""
    result = {}
    for key, val in sess.items():
        if key == "db_connector":
            result[key] = None
            continue
        if isinstance(val, dict):
            inner = {}
            for k, v in val.items():
                if hasattr(v, "model_dump"):
                    inner[k] = v.model_dump()
                elif isinstance(v, dict):
                    nested = {}
                    for nk, nv in v.items():
                        if hasattr(nv, "model_dump"):
                            nested[nk] = nv.model_dump()
                        else:
                            nested[nk] = nv
                    inner[k] = nested
                else:
                    inner[k] = v
            result[key] = inner
        else:
            result[key] = val
    return result


def sse(payload: dict) -> str:
    return f"data: {json.dumps(payload, default=str)}\n\n"


def _superset_client_from_session(sess: dict):
    from tools.superset_api import SupersetClient

    superset_cfg = sess["superset"]
    client = SupersetClient(
        base_url=superset_cfg["url"] or config.SUPERSET_URL,
        username=superset_cfg["username"] or config.SUPERSET_USERNAME,
        password=superset_cfg["password"] or config.SUPERSET_PASSWORD,
        session_cookie=superset_cfg.get("session_cookie") or None,
        csrf_token=superset_cfg.get("csrf_token") or None,
    )
    client.authenticate()
    return client


def _created_by_from_session(sess: dict) -> str:
    superset_cfg = sess.get("superset", {})
    return superset_cfg.get("username") or config.SUPERSET_USERNAME or "unknown"


def _extract_filters_from_metadata(json_metadata: dict) -> list[dict]:
    filters = []
    for item in json_metadata.get("native_filter_configuration", []) or []:
        targets = item.get("targets", []) or []
        first_target = targets[0] if targets else {}
        column = first_target.get("column", {})
        filters.append(
            {
                "id": item.get("id", ""),
                "name": item.get("name", ""),
                "filter_type": item.get("filterType", ""),
                "column": column.get("name", "") if isinstance(column, dict) else "",
            }
        )
    return filters


def _snapshot_charts_from_full_state(state: dict) -> list[dict]:
    charts = []
    for chart in state.get("slices", []) or []:
        charts.append(
            {
                "id": chart.get("id") or 0,
                "title": chart.get("title", ""),
                "viz_type": chart.get("viz_type", ""),
                "dataset_id": chart.get("datasource_id") or 0,
                "params": chart.get("params", {}),
            }
        )
    return charts


def _save_full_state_snapshot(
    *,
    version_history,
    dashboard_id: int,
    state: dict,
    dashboard_url: str,
    requirements_prompt: str,
    dataset_name: str,
    created_by: str,
    change_summary: str,
) -> int:
    return version_history.save_snapshot(
        dashboard_id=dashboard_id,
        dashboard_title=state.get("title", ""),
        dashboard_url=dashboard_url,
        requirements_prompt=requirements_prompt,
        dataset_name=dataset_name,
        position_json=state.get("position_json", {}),
        charts=_snapshot_charts_from_full_state(state),
        filters=_extract_filters_from_metadata(state.get("json_metadata", {})),
        created_by=created_by,
        change_summary=change_summary,
    )


def _summarize_dashboard_change(
    prev_snapshot: dict | None,
    new_chart_titles: list[str],
    fallback: str,
) -> str:
    if not prev_snapshot:
        return fallback
    try:
        from tools.llm_client import chat

        prev_titles = [
            c.get("title", "")
            for c in prev_snapshot.get("charts", [])
            if isinstance(c, dict)
        ]
        prompt = (
            "Previous chart titles:\n"
            f"{json.dumps(prev_titles)}\n\n"
            "New chart titles:\n"
            f"{json.dumps(new_chart_titles)}\n\n"
            "Write one short dashboard version change summary under 120 characters. "
            "Mention added or removed charts when obvious. Return plain text only."
        )
        summary = chat(
            "You summarize dashboard configuration changes concisely.",
            prompt,
            max_tokens=80,
            temperature=0,
        ).strip()
        return summary[:180] if summary else fallback
    except Exception:
        return fallback


# ── FastAPI app ────────────────────────────────────────────────────────────────

app = FastAPI(title="DATA VIZ API")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ── Request / Response models ─────────────────────────────────────────────────

class DbConfig(BaseModel):
    type: str = "postgresql"
    host: str = "localhost"
    port: int = 5432
    database: str = ""
    username: str = ""
    password: str = ""


class SupersetConfig(BaseModel):
    url: str = "http://localhost:8088"
    username: str = "admin"
    password: str = ""
    session_cookie: str = ""
    csrf_token: str = ""


class SessionConfigUpdate(BaseModel):
    db: DbConfig | None = None
    superset: SupersetConfig | None = None
    llm_model: str | None = None


class ProfileTableBody(BaseModel):
    table_name: str


class AddTableBody(BaseModel):
    table_name: str
    selected_columns: list[str]


class ConfirmPhase2Body(BaseModel):
    edited_sql: str


class CsvQueryBody(BaseModel):
    question: str
    history: list[dict] = []


class CsvChartBody(BaseModel):
    id: str
    title: str
    chart_type: str
    columns: list[str]
    rows: list[dict]
    x_col: str
    y_col: str
    question: str
    added_at: str


class ChartDescriptionSpec(BaseModel):
    id: int
    title: str
    viz_type: str
    metrics: list[Any]
    groupby: list[str]
    time_column: str | None = None
    time_grain: str | None = None


class GenerateDescriptionsBody(BaseModel):
    chart_ids: list[int]
    chart_specs: list[ChartDescriptionSpec]
    dataset_name: str
    dashboard_title: str


def _is_readonly_select_sql(sql: str) -> bool:
    stripped = sql.strip().rstrip(";")
    if not stripped:
        return False
    if ";" in stripped:
        return False

    compact = " ".join(stripped.split())
    if not re.match(r"^(select|with)\b", compact, re.IGNORECASE):
        return False

    forbidden = (
        "insert", "update", "delete", "drop", "alter", "create", "replace",
        "truncate", "merge", "attach", "detach", "copy", "export", "import",
        "call", "pragma", "set", "install", "load",
    )
    return not re.search(r"\b(" + "|".join(forbidden) + r")\b", compact, re.IGNORECASE)


# ── Session endpoints ─────────────────────────────────────────────────────────

@app.post("/api/session")
async def create_session():
    sid = str(uuid.uuid4())
    sessions[sid] = new_session()
    return {"session_id": sid}


@app.get("/api/config/defaults")
async def get_defaults():
    return {
        "superset_url": config.SUPERSET_URL,
        "superset_username": config.SUPERSET_USERNAME,
        "llm_model": config.LLM_MODEL,
    }


@app.get("/api/sessions/{sid}/state")
async def get_state(sid: str):
    sess = get_session(sid)
    return serialize_session(sess)


@app.put("/api/sessions/{sid}/config")
async def update_config(sid: str, body: SessionConfigUpdate):
    sess = get_session(sid)
    if body.db is not None:
        sess["db"] = body.db.model_dump()
        # Reset connector when config changes
        sess["db_connector"] = None
    if body.superset is not None:
        sess["superset"] = body.superset.model_dump()
    if body.llm_model is not None:
        sess["llm_model"] = body.llm_model
    return {"ok": True}


@app.post("/api/sessions/{sid}/db/test")
async def test_db_connection(sid: str):
    sess = get_session(sid)
    db = sess["db"]
    try:
        from tools.db_connector import DBConnector
        connector = DBConnector(
            db_type=db["type"],
            host=db["host"],
            port=int(db["port"]),
            database=db["database"],
            username=db["username"],
            password=db["password"],
        )
        ok, message = connector.test_connection()
        if ok:
            sess["db_connector"] = connector
        return {"ok": ok, "message": message}
    except Exception as e:
        return {"ok": False, "message": str(e)}


# ── Phase 1 endpoints ─────────────────────────────────────────────────────────

@app.get("/api/sessions/{sid}/phase1/explore")
async def phase1_explore(sid: str, prompt: str = Query(...)):
    sess = get_session(sid)
    queue: asyncio.Queue = asyncio.Queue()
    loop = asyncio.get_event_loop()

    def worker():
        try:
            if sess.get("llm_model"):
                config.LLM_MODEL = sess["llm_model"]

            db_connector = sess.get("db_connector")
            if db_connector is None:
                loop.call_soon_threadsafe(
                    queue.put_nowait,
                    {"type": "error", "message": "No database connection. Test connection first."},
                )
                return

            append_audit(sess, 1, "prompt_entered", "Prompt entered", prompt, status="info")

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "progress", "message": "Getting all tables from database...", "step": 1, "total": 3},
            )

            from agents.schema_explorer import SchemaExplorer
            explorer = SchemaExplorer()
            profiled_tables = explorer.run(prompt, db_connector)

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "progress", "message": f"Profiled {len(profiled_tables)} tables. Analyzing context...", "step": 2, "total": 3},
            )

            from agents.context_analyst import ContextAnalyst
            analyst = ContextAnalyst()
            schema_map = analyst.run(prompt, profiled_tables)

            append_audit(
                sess, 1, "tables_discovered", "Tables discovered",
                f"Found {len(schema_map.all_tables)} tables in database",
                data={"all_tables": schema_map.all_tables, "count": len(schema_map.all_tables)},
                status="info",
            )

            selected_names_set = {t.table_name for t in schema_map.profiled_tables}
            append_audit(
                sess, 1, "tables_selected",
                f"Agent selected {len(schema_map.profiled_tables)} tables",
                schema_map.agent_reasoning,
                data={
                    "selected": [t.table_name for t in schema_map.profiled_tables],
                    "excluded": [t for t in schema_map.all_tables if t not in selected_names_set],
                    "suggested_primary": schema_map.suggested_primary,
                    "suggested_joins": schema_map.suggested_joins,
                },
                status="success",
            )

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "progress", "message": "Profiling excluded tables...", "step": 3, "total": 3},
            )

            # Profile excluded tables
            selected_names = {t.table_name for t in schema_map.profiled_tables}
            all_tables = schema_map.all_tables
            excluded_tables: dict[str, Any] = {}

            for tname in all_tables:
                if tname not in selected_names:
                    try:
                        profile = db_connector.profile_table(tname)
                        excluded_tables[tname] = {
                            "profiled": True,
                            "profile": profile.model_dump(),
                            "selected_columns": [],
                            "added": False,
                            "error": None,
                        }
                    except Exception as e:
                        excluded_tables[tname] = {
                            "profiled": False,
                            "profile": None,
                            "selected_columns": [],
                            "added": False,
                            "error": str(e),
                        }

            sess["phase1"]["schema_map"] = schema_map
            sess["phase1"]["business_prompt"] = prompt
            sess["phase1"]["excluded_tables"] = excluded_tables
            sess["phase1"]["confirmed"] = False

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {
                    "type": "done",
                    "data": {
                        "schema_map": schema_map.model_dump(),
                        "excluded_tables": excluded_tables,
                    },
                },
            )
        except Exception as e:
            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "error", "message": str(e)},
            )

    loop.run_in_executor(None, worker)

    async def generate() -> AsyncGenerator[str, None]:
        while True:
            try:
                item = await asyncio.wait_for(queue.get(), timeout=300)
            except asyncio.TimeoutError:
                yield sse({"type": "error", "message": "Timeout waiting for agent response"})
                break
            yield sse(item)
            if item["type"] in ("done", "error"):
                break

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.post("/api/sessions/{sid}/phase1/profile-table")
async def phase1_profile_table(sid: str, body: ProfileTableBody):
    sess = get_session(sid)
    db_connector = sess.get("db_connector")
    if db_connector is None:
        raise HTTPException(status_code=400, detail="No database connection")
    try:
        profile = db_connector.profile_table(body.table_name)
        excluded = sess["phase1"]["excluded_tables"]
        if body.table_name not in excluded:
            excluded[body.table_name] = {
                "profiled": False,
                "profile": None,
                "selected_columns": [],
                "added": False,
                "error": None,
            }
        excluded[body.table_name]["profiled"] = True
        excluded[body.table_name]["profile"] = profile.model_dump()
        excluded[body.table_name]["error"] = None
        return profile.model_dump()
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/sessions/{sid}/phase1/add-table")
async def phase1_add_table(sid: str, body: AddTableBody):
    sess = get_session(sid)
    excluded = sess["phase1"]["excluded_tables"]
    schema_map = sess["phase1"].get("schema_map")
    if schema_map is None:
        raise HTTPException(status_code=400, detail="No schema map. Run explore first.")

    entry = excluded.get(body.table_name)
    if entry is None or not entry.get("profiled") or entry.get("profile") is None:
        raise HTTPException(status_code=400, detail="Table not profiled yet")

    from models.schemas import ColumnProfile, TableProfile

    raw_profile = entry["profile"]
    all_cols = raw_profile.get("columns", [])
    if body.selected_columns:
        filtered_cols = [c for c in all_cols if c["column_name"] in body.selected_columns]
    else:
        filtered_cols = all_cols

    tp = TableProfile(
        table_name=raw_profile["table_name"],
        row_count=raw_profile["row_count"],
        columns=[ColumnProfile(**c) for c in filtered_cols],
        sample_rows=raw_profile.get("sample_rows", []),
    )

    existing_names = {t.table_name for t in schema_map.profiled_tables}
    if tp.table_name not in existing_names:
        schema_map.profiled_tables.append(tp)
        if tp.table_name not in schema_map.all_tables:
            schema_map.all_tables.append(tp.table_name)

    excluded[body.table_name]["added"] = True
    excluded[body.table_name]["selected_columns"] = body.selected_columns

    col_list = ", ".join(body.selected_columns) if body.selected_columns else "all columns"
    append_audit(
        sess, 1, "table_manually_added",
        f"User added table: {body.table_name}",
        f"Columns selected: {col_list}",
        data={"table_name": body.table_name, "columns": body.selected_columns},
        status="info",
    )

    return {"ok": True}


@app.post("/api/sessions/{sid}/phase1/confirm")
async def phase1_confirm(sid: str):
    sess = get_session(sid)
    if sess["phase1"].get("schema_map") is None:
        raise HTTPException(status_code=400, detail="No schema map to confirm")
    sess["phase1"]["confirmed"] = True
    schema_map = sess["phase1"]["schema_map"]
    n = len(schema_map.profiled_tables) if hasattr(schema_map, "profiled_tables") else 0
    append_audit(
        sess, 1, "phase1_confirmed", "Schema confirmed",
        f"Proceeding with {n} tables",
        status="success",
    )
    return {"ok": True}


@app.get("/api/sessions/{sid}/phase1/erd-data")
async def phase1_erd_data(sid: str):
    sess = get_session(sid)
    schema_map_raw = sess["phase1"].get("schema_map")
    if schema_map_raw is None:
        raise HTTPException(status_code=404, detail="No schema map yet")

    schema_map = (
        schema_map_raw.model_dump()
        if hasattr(schema_map_raw, "model_dump")
        else schema_map_raw
    )

    selected_nodes = []
    for table in schema_map.get("profiled_tables", []):
        table_data = table.model_dump() if hasattr(table, "model_dump") else table
        columns = []
        for col in table_data.get("columns", []):
            col_data = col.model_dump() if hasattr(col, "model_dump") else col
            columns.append(
                {
                    "name": col_data["column_name"],
                    "type": col_data["data_type"],
                    "is_pk": col_data["is_likely_pk"],
                    "is_fk": col_data["is_likely_fk"],
                    "is_date": col_data["is_likely_date"],
                    "null_pct": col_data["null_pct"],
                }
            )
        selected_nodes.append(
            {
                "table_name": table_data["table_name"],
                "row_count": table_data["row_count"],
                "columns": columns,
                "is_primary": table_data["table_name"] == schema_map.get("suggested_primary", ""),
                "is_selected": True,
            }
        )

    selected_names = {node["table_name"] for node in selected_nodes}
    excluded = sess["phase1"].get("excluded_tables", {})
    excluded_nodes = []
    for table_name in schema_map.get("all_tables", []):
        if table_name in selected_names:
            continue
        entry = excluded.get(table_name, {})
        profile = entry.get("profile") if isinstance(entry, dict) else None
        columns = []
        row_count = None
        if profile:
            for col in profile.get("columns", []):
                columns.append(
                    {
                        "name": col["column_name"],
                        "type": col["data_type"],
                        "is_pk": col["is_likely_pk"],
                        "is_fk": col["is_likely_fk"],
                        "is_date": col["is_likely_date"],
                        "null_pct": col["null_pct"],
                    }
                )
            row_count = profile.get("row_count")
        excluded_nodes.append(
            {
                "table_name": table_name,
                "row_count": row_count,
                "columns": columns,
                "is_primary": False,
                "is_selected": False,
            }
        )

    edges = []
    for join_str in schema_map.get("suggested_joins", []):
        if str(join_str).strip().startswith("--"):
            continue
        match = re.search(
            r"JOIN\s+([\w.]+)\s+ON\s+([\w.]+)\.(\w+)\s*=\s*([\w.]+)\.(\w+)",
            str(join_str),
            re.IGNORECASE,
        )
        if match:
            edges.append(
                {
                    "join_table": match.group(1),
                    "source_table": match.group(2),
                    "source_col": match.group(3),
                    "target_table": match.group(4),
                    "target_col": match.group(5),
                    "join_string": str(join_str),
                }
            )

    return {
        "selected_nodes": selected_nodes,
        "excluded_nodes": excluded_nodes,
        "edges": edges,
        "suggested_primary": schema_map.get("suggested_primary", ""),
    }


# ── Phase 2 endpoints ─────────────────────────────────────────────────────────

@app.get("/api/sessions/{sid}/phase2/generate")
async def phase2_generate(sid: str):
    sess = get_session(sid)
    if not sess["phase1"].get("confirmed"):
        raise HTTPException(status_code=400, detail="Phase 1 not confirmed")

    queue: asyncio.Queue = asyncio.Queue()
    loop = asyncio.get_event_loop()

    def worker():
        try:
            if sess.get("llm_model"):
                config.LLM_MODEL = sess["llm_model"]

            schema_map = sess["phase1"].get("schema_map")
            db_connector = sess.get("db_connector")
            business_prompt = sess["phase1"].get("business_prompt", "")

            if schema_map is None:
                loop.call_soon_threadsafe(
                    queue.put_nowait,
                    {"type": "error", "message": "No schema map found"},
                )
                return

            MAX_RETRIES = 2
            last_error = ""
            query_plan = None

            for attempt in range(MAX_RETRIES + 1):
                loop.call_soon_threadsafe(
                    queue.put_nowait,
                    {
                        "type": "progress",
                        "message": f"Writing JOIN query{'(retry ' + str(attempt) + ')' if attempt > 0 else ''}...",
                        "step": 1,
                        "total": 2,
                    },
                )

                if attempt > 0 and last_error:
                    append_audit(
                        sess, 2, "query_retry", "Query retry triggered",
                        last_error,
                        data={"error": last_error, "attempt": attempt},
                        status="warning",
                    )

                from agents.query_architect import QueryArchitect
                architect = QueryArchitect()

                prompt_with_error = business_prompt
                if last_error:
                    prompt_with_error += f"\n\nPrevious SQL failed with error: {last_error}\nPlease fix the SQL."

                query_plan = architect.run(prompt_with_error, schema_map)

                append_audit(
                    sess, 2, "query_generated",
                    f"Query generated (attempt {attempt + 1})",
                    query_plan.grain_description or "Query generated",
                    data={
                        "attempt": attempt + 1,
                        "sql_preview": query_plan.sql[:200] + ("..." if len(query_plan.sql) > 200 else ""),
                        "calculated_columns": [c["name"] if isinstance(c, dict) else c.name for c in query_plan.calculated_columns],
                    },
                    status="info",
                )

                if db_connector is None:
                    break

                loop.call_soon_threadsafe(
                    queue.put_nowait,
                    {"type": "progress", "message": "Running QA checks against database...", "step": 2, "total": 2},
                )

                from agents.dataset_qa import DatasetQA
                qa_agent = DatasetQA()
                qa_report = qa_agent.run(query_plan, db_connector)

                append_audit(
                    sess, 2, "qa_result",
                    f"QA {'passed' if qa_report.passed else 'failed'}",
                    "; ".join(qa_report.issues) if qa_report.issues else "No issues found",
                    data={
                        "passed": qa_report.passed,
                        "row_count": qa_report.row_count,
                        "duplicate_count": qa_report.duplicate_row_count,
                        "issues": qa_report.issues,
                        "suggestions": qa_report.suggestions,
                    },
                    status="success" if qa_report.passed else "error",
                )

                sess["phase2"]["query_plan"] = query_plan
                sess["phase2"]["qa_report"] = qa_report
                sess["phase2"]["edited_sql"] = query_plan.sql
                sess["phase2"]["confirmed"] = False

                if qa_report.passed or attempt >= MAX_RETRIES:
                    loop.call_soon_threadsafe(
                        queue.put_nowait,
                        {
                            "type": "done",
                            "data": {
                                "query_plan": query_plan.model_dump(),
                                "qa_report": qa_report.model_dump(),
                            },
                        },
                    )
                    return
                else:
                    sql_errors = [i for i in qa_report.issues if "execution failed" in i.lower() or "syntax" in i.lower()]
                    if sql_errors:
                        last_error = "; ".join(sql_errors)
                    else:
                        # Non-SQL QA issues (duplicates, fan-out) — send done with current results
                        loop.call_soon_threadsafe(
                            queue.put_nowait,
                            {
                                "type": "done",
                                "data": {
                                    "query_plan": query_plan.model_dump(),
                                    "qa_report": qa_report.model_dump(),
                                },
                            },
                        )
                        return

            # No db_connector path
            if query_plan is not None and db_connector is None:
                sess["phase2"]["query_plan"] = query_plan
                sess["phase2"]["edited_sql"] = query_plan.sql
                sess["phase2"]["confirmed"] = False
                loop.call_soon_threadsafe(
                    queue.put_nowait,
                    {
                        "type": "done",
                        "data": {
                            "query_plan": query_plan.model_dump(),
                            "qa_report": None,
                        },
                    },
                )

        except Exception as e:
            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "error", "message": str(e)},
            )

    loop.run_in_executor(None, worker)

    async def generate() -> AsyncGenerator[str, None]:
        while True:
            try:
                item = await asyncio.wait_for(queue.get(), timeout=300)
            except asyncio.TimeoutError:
                yield sse({"type": "error", "message": "Timeout"})
                break
            yield sse(item)
            if item["type"] in ("done", "error"):
                break

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.post("/api/sessions/{sid}/phase2/confirm")
async def phase2_confirm(sid: str, body: ConfirmPhase2Body):
    sess = get_session(sid)
    if sess["phase2"].get("query_plan") is None:
        raise HTTPException(status_code=400, detail="No query plan to confirm")
    sess["phase2"]["confirmed"] = True
    sess["phase2"]["edited_sql"] = body.edited_sql
    # Pre-fill phase3 dataset name
    qp = sess["phase2"]["query_plan"]
    dataset_name = ""
    if qp and hasattr(qp, "dataset_name_suggestion"):
        dataset_name = qp.dataset_name_suggestion
        sess["phase3"]["dataset_name"] = dataset_name
    append_audit(
        sess, 2, "phase2_confirmed", "Query confirmed",
        f"Dataset name: {dataset_name}",
        data={"dataset_name_suggestion": dataset_name},
        status="success",
    )
    return {"ok": True}


# ── Phase 3 endpoints ─────────────────────────────────────────────────────────

@app.get("/api/sessions/{sid}/phase3/plan")
async def phase3_plan(
    sid: str,
    dataset_name: str = Query(...),
    dashboard_title: str = Query(...),
    requirements: str = Query(...),
):
    sess = get_session(sid)
    queue: asyncio.Queue = asyncio.Queue()
    loop = asyncio.get_event_loop()

    def worker():
        try:
            if sess.get("llm_model"):
                config.LLM_MODEL = sess["llm_model"]

            superset_cfg = sess["superset"]
            _superset_url = superset_cfg["url"] or config.SUPERSET_URL
            logger.debug("Phase3 plan using Superset URL: %s", _superset_url)

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "progress", "message": "Authenticating with Superset...", "step": 1, "total": 5},
            )
            from tools.superset_api import SupersetClient
            _session_cookie = superset_cfg.get("session_cookie", "")
            _csrf_token = superset_cfg.get("csrf_token", "")
            client = SupersetClient(
                base_url=_superset_url,
                username=superset_cfg["username"] or config.SUPERSET_USERNAME,
                password=superset_cfg["password"] or config.SUPERSET_PASSWORD,
                session_cookie=_session_cookie or None,
                csrf_token=_csrf_token or None,
            )
            client.authenticate()

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "progress", "message": "Fetching dataset info...", "step": 2, "total": 5},
            )
            dataset_info = client.get_dataset_by_name(dataset_name)

            append_audit(
                sess, 3, "dataset_loaded",
                f"Dataset loaded: {dataset_info.name}",
                f"{len(dataset_info.columns)} columns",
                data={"dataset_name": dataset_info.name, "column_count": len(dataset_info.columns)},
                status="info",
            )

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "progress", "message": "Sampling column values...", "step": 3, "total": 5},
            )
            from tools.column_sampler import ColumnSampler
            sampler = ColumnSampler(client)
            dataset_info = sampler.enrich_columns(dataset_info)

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "progress", "message": "Parsing requirements...", "step": 4, "total": 5},
            )
            from agents.requirements_parser import RequirementsParser
            parser = RequirementsParser()
            parsed_reqs = parser.run(requirements, dataset_info)

            flagged = parsed_reqs.get("flagged", []) if isinstance(parsed_reqs, dict) else []
            charts_list = parsed_reqs.get("charts", []) if isinstance(parsed_reqs, dict) else []
            req_detail = f"{len(charts_list)} chart intents extracted"
            if flagged:
                req_detail += f", {len(flagged)} flagged"
            append_audit(
                sess, 3, "requirements_parsed", "Requirements parsed",
                req_detail,
                data={
                    "chart_count": len(charts_list),
                    "flagged": flagged,
                    "chart_intents": [c.get("intent", "") for c in charts_list if isinstance(c, dict)],
                },
                status="warning" if flagged else "success",
            )

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "progress", "message": "Planning charts...", "step": 5, "total": 5},
            )
            from agents.chart_strategist import ChartStrategist
            from tools.catalogue import CatalogueManager
            catalogue = CatalogueManager()
            strategist = ChartStrategist()
            dashboard_plan = strategist.run(parsed_reqs, dataset_info, catalogue, dashboard_title)

            plan_detail = ", ".join(
                f"{c.viz_type}: {c.title}" for c in dashboard_plan.charts
            )
            append_audit(
                sess, 3, "dashboard_planned",
                f"Dashboard planned: {len(dashboard_plan.charts)} charts",
                plan_detail,
                data={
                    "charts": [
                        {"title": c.title, "viz_type": c.viz_type,
                         "reasoning": c.reasoning, "width": c.width}
                        for c in dashboard_plan.charts
                    ],
                    "filters": [f.column_name for f in dashboard_plan.filters],
                },
                status="success",
            )

            sess["phase3"]["dataset_name"] = dataset_name
            sess["phase3"]["dashboard_title"] = dashboard_title
            sess["phase3"]["requirements"] = requirements
            sess["phase3"]["dataset_info"] = dataset_info
            sess["phase3"]["parsed_reqs"] = parsed_reqs
            sess["phase3"]["dashboard_plan"] = dashboard_plan
            sess["phase3"]["plan_ready"] = True

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {
                    "type": "done",
                    "data": {
                        "dataset_info": dataset_info.model_dump(),
                        "parsed_reqs": parsed_reqs,
                        "dashboard_plan": dashboard_plan.model_dump(),
                    },
                },
            )
        except Exception as e:
            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "error", "message": str(e)},
            )

    loop.run_in_executor(None, worker)

    async def generate() -> AsyncGenerator[str, None]:
        while True:
            try:
                item = await asyncio.wait_for(queue.get(), timeout=300)
            except asyncio.TimeoutError:
                yield sse({"type": "error", "message": "Timeout"})
                break
            yield sse(item)
            if item["type"] in ("done", "error"):
                break

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


@app.get("/api/sessions/{sid}/phase3/build")
async def phase3_build(
    sid: str,
    dashboard_id: str | None = Query(default=None),
    dry_run: bool = Query(default=False),
):
    sess = get_session(sid)

    queue: asyncio.Queue = asyncio.Queue()
    loop = asyncio.get_event_loop()

    def worker():
        try:
            if not sess["phase3"].get("plan_ready"):
                loop.call_soon_threadsafe(
                    queue.put_nowait,
                    {"type": "error", "message": "No dashboard plan found. Please run 'Plan Dashboard' first."},
                )
                return

            if sess.get("llm_model"):
                config.LLM_MODEL = sess["llm_model"]

            if dry_run:
                plan = sess["phase3"]["dashboard_plan"]
                loop.call_soon_threadsafe(
                    queue.put_nowait,
                    {
                        "type": "done",
                        "data": {
                            "dry_run": True,
                            "dashboard_plan": plan.model_dump() if hasattr(plan, "model_dump") else plan,
                            "dashboard_url": None,
                            "qa_report": None,
                            "chart_ids": [],
                            "chart_count": 0,
                            "filter_count": 0,
                            "chart_actions": [],
                        },
                    },
                )
                return

            superset_cfg = sess["superset"]
            dataset_info = sess["phase3"]["dataset_info"]
            dashboard_plan = sess["phase3"]["dashboard_plan"]

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "progress", "message": "Authenticating with Superset...", "step": 1, "total": 7},
            )
            from tools.superset_api import SupersetClient, build_chart_params, build_position_json
            _session_cookie = superset_cfg.get("session_cookie", "")
            _csrf_token = superset_cfg.get("csrf_token", "")
            client = SupersetClient(
                base_url=superset_cfg["url"] or config.SUPERSET_URL,
                username=superset_cfg["username"] or config.SUPERSET_USERNAME,
                password=superset_cfg["password"] or config.SUPERSET_PASSWORD,
                session_cookie=_session_cookie or None,
                csrf_token=_csrf_token or None,
            )
            client.authenticate()

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "progress", "message": "Fetching existing charts...", "step": 2, "total": 7},
            )
            existing_charts = client.get_charts_for_dataset(dataset_info.id)

            chart_ids: list[int] = []
            chart_actions: list[tuple[int, str]] = []
            total_charts = len(dashboard_plan.charts)

            for i, chart_spec in enumerate(dashboard_plan.charts):
                loop.call_soon_threadsafe(
                    queue.put_nowait,
                    {
                        "type": "progress",
                        "message": f"Creating chart {i + 1}/{total_charts}: {chart_spec.title}...",
                        "step": 3,
                        "total": 7,
                    },
                )
                chart_id, action = client.upsert_chart(dataset_info.id, chart_spec, existing_charts)
                chart_ids.append(int(chart_id))
                chart_actions.append((int(chart_id), action))
                append_audit(
                    sess, 3, "chart_created",
                    f"Chart created: {chart_spec.title}",
                    f"{chart_spec.viz_type} · {action}",
                    data={"chart_id": chart_id, "action": action},
                    status="success",
                )
            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "progress", "message": "Building dashboard layout...", "step": 4, "total": 7},
            )
            position_json = build_position_json(chart_ids, dashboard_plan.charts)

            # Sanity-check: every chart ID in position_json must match the collected IDs
            in_layout = {
                v["meta"]["chartId"]
                for v in position_json.values()
                if isinstance(v, dict) and v.get("type") == "CHART"
            }
            collected = set(chart_ids)
            if in_layout != collected:
                raise ValueError(
                    f"position_json ID mismatch — "
                    f"collected={sorted(collected)}, "
                    f"in_layout={sorted(in_layout)}"
                )

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "progress", "message": "Creating/Updating dashboard...", "step": 5, "total": 7},
            )

            try:
                did = int(dashboard_id) if dashboard_id else None
            except ValueError:
                raise ValueError(f"Invalid dashboard ID '{dashboard_id}' — must be a number")

            is_update = did is not None
            from tools.version_history import DashboardVersionHistory
            version_history = DashboardVersionHistory()
            created_by = _created_by_from_session(sess)

            if did:
                # First app-managed update for an existing dashboard: preserve the
                # old Superset state before overwriting it so rollback is possible.
                if not version_history.load_index(did):
                    previous_state = client.get_dashboard_full_state(did)
                    if previous_state:
                        _save_full_state_snapshot(
                            version_history=version_history,
                            dashboard_id=did,
                            state=previous_state,
                            dashboard_url=f"{client.base_url}/superset/dashboard/{did}",
                            requirements_prompt=sess["phase3"].get("requirements", ""),
                            dataset_name=sess["phase3"].get("dataset_name", ""),
                            created_by=created_by,
                            change_summary="Initial version",
                        )
                dashboard_url = client.update_dashboard(did, chart_ids, position_json)
                dash_action = "updated"
            else:
                did, dashboard_url = client.create_dashboard(
                    dashboard_plan.dashboard_title, chart_ids, position_json
                )
                dash_action = "created"

            append_audit(
                sess, 3, "dashboard_created",
                f"Dashboard {dash_action} in Superset",
                f"{len(chart_ids)} charts · {len(dashboard_plan.filters)} filters",
                data={
                    "dashboard_id": did,
                    "dashboard_url": dashboard_url,
                    "chart_count": len(chart_ids),
                    "filter_count": len(dashboard_plan.filters),
                },
                status="success",
            )

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "progress", "message": "Configuring filters...", "step": 6, "total": 7},
            )
            client.set_dashboard_filters(did, dashboard_plan.filters, dataset_info.id)

            saved_version = -1
            try:
                prev_index = version_history.load_index(did) if is_update else None
                prev_snapshot = None
                if prev_index and prev_index.get("versions"):
                    prev_snapshot = version_history.load_snapshot(
                        did,
                        prev_index["versions"][-1]["version"],
                    )

                if not is_update or not prev_snapshot:
                    change_summary = "Initial version"
                else:
                    prev_charts = {
                        c.get("title", "")
                        for c in prev_snapshot.get("charts", [])
                        if isinstance(c, dict)
                    }
                    new_charts = {c.title for c in dashboard_plan.charts}
                    added = sorted(new_charts - prev_charts)
                    removed = sorted(prev_charts - new_charts)
                    parts = []
                    if added:
                        parts.append(f"Added: {', '.join(added)}")
                    if removed:
                        parts.append(f"Removed: {', '.join(removed)}")
                    fallback_summary = "; ".join(parts) if parts else "Updated chart configurations"
                    change_summary = _summarize_dashboard_change(
                        prev_snapshot,
                        [c.title for c in dashboard_plan.charts],
                        fallback_summary,
                    )

                snapshot_charts = [
                    {
                        "id": cid,
                        "title": spec.title,
                        "viz_type": spec.viz_type,
                        "dataset_id": dataset_info.id,
                        "params": build_chart_params(spec),
                    }
                    for cid, spec in zip(chart_ids, dashboard_plan.charts)
                ]
                snapshot_filters = [
                    {
                        "id": f"filter_{i}",
                        "name": filt.label,
                        "filter_type": filt.filter_type,
                        "column": filt.column_name,
                    }
                    for i, filt in enumerate(dashboard_plan.filters)
                ]
                saved_version = version_history.save_snapshot(
                    dashboard_id=did,
                    dashboard_title=dashboard_plan.dashboard_title,
                    dashboard_url=dashboard_url,
                    requirements_prompt=sess["phase3"].get("requirements", ""),
                    dataset_name=sess["phase3"].get("dataset_name", ""),
                    position_json=position_json,
                    charts=snapshot_charts,
                    filters=snapshot_filters,
                    created_by=created_by,
                    change_summary=change_summary,
                )
                message = (
                    f"Version {saved_version} saved to history"
                    if saved_version > 0
                    else "Version history save failed (non-fatal)"
                )
                loop.call_soon_threadsafe(
                    queue.put_nowait,
                    {
                        "type": "progress",
                        "message": message,
                        "step": 6,
                        "total": 7,
                    },
                )
            except Exception as e:
                loop.call_soon_threadsafe(
                    queue.put_nowait,
                    {
                        "type": "progress",
                        "message": f"Version history save failed (non-fatal): {e}",
                        "step": 6,
                        "total": 7,
                    },
                )

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "progress", "message": "Running QA review...", "step": 7, "total": 7},
            )
            from agents.qa_reviewer import QAReviewer
            qa = QAReviewer()
            qa_report = qa.run(dashboard_plan, dataset_info, chart_actions, client)

            append_audit(
                sess, 3, "qa_review",
                f"QA {'passed' if qa_report.passed else 'failed'}",
                "; ".join(qa_report.issues) if qa_report.issues else "All checks passed",
                data={"passed": qa_report.passed, "issues": qa_report.issues},
                status="success" if qa_report.passed else "warning",
            )

            append_audit(
                sess, 3, "dashboard_live",
                "Dashboard live",
                dashboard_url,
                data={"url": dashboard_url},
                status="success",
            )

            # Persist audit log to disk
            try:
                runs_dir = Path("runs")
                runs_dir.mkdir(exist_ok=True)
                ts = datetime.now().strftime("%Y%m%d_%H%M%S")
                audit_file = runs_dir / f"audit_{ts}.json"
                with open(audit_file, "w") as _f:
                    json.dump(
                        {"session_id": sid, "created_at": ts, "entries": sess["audit_log"]},
                        _f,
                        indent=2,
                        default=str,
                    )
            except Exception:
                pass

            sess["phase3"]["dashboard_url"] = dashboard_url
            sess["phase3"]["dashboard_id"] = did
            sess["phase3"]["chart_ids"] = chart_ids

            # Update catalogue
            try:
                from tools.catalogue import CatalogueManager
                from models.schemas import CatalogueEntry
                catalogue = CatalogueManager()
                new_entries = []
                for chart_spec, (chart_id, action) in zip(dashboard_plan.charts, chart_actions):
                    if action == "created":
                        entry = CatalogueEntry(
                            client_hint=sess["phase3"].get("dashboard_title", ""),
                            intent=chart_spec.title,
                            viz_type=chart_spec.viz_type,
                            metric_columns=[
                                (m.get("column") or {}).get("column_name", "")
                                for m in chart_spec.metrics
                                if m.get("expressionType") == "SIMPLE"
                            ],
                            dimension_columns=chart_spec.groupby,
                            time_column=chart_spec.time_column,
                            worked_well=qa_report.passed,
                            notes="; ".join(qa_report.suggestions[:2]),
                        )
                        new_entries.append(entry)
                if new_entries:
                    catalogue.append(new_entries)
            except Exception:
                pass

            loop.call_soon_threadsafe(
                queue.put_nowait,
                {
                    "type": "done",
                    "data": {
                        "dashboard_id": did,
                        "dashboard_url": dashboard_url,
                        "qa_report": qa_report.model_dump(),
                        "chart_ids": chart_ids,
                        "chart_count": len(chart_ids),
                        "filter_count": len(dashboard_plan.filters),
                        "saved_version": saved_version,
                        "chart_actions": [
                            {"chart_id": cid, "action": act} for cid, act in chart_actions
                        ],
                    },
                },
            )
        except Exception as e:
            loop.call_soon_threadsafe(
                queue.put_nowait,
                {"type": "error", "message": str(e)},
            )

    loop.run_in_executor(None, worker)

    async def generate() -> AsyncGenerator[str, None]:
        while True:
            try:
                item = await asyncio.wait_for(queue.get(), timeout=300)
            except asyncio.TimeoutError:
                yield sse({"type": "error", "message": "Timeout"})
                break
            yield sse(item)
            if item["type"] in ("done", "error"):
                break

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={"Cache-Control": "no-cache", "X-Accel-Buffering": "no"},
    )


# ── Phase 3 extra endpoints ───────────────────────────────────────────────────


class UpdatePlanBody(BaseModel):
    plan: Any


@app.get("/api/sessions/{sid}/phase3/chart-preview")
async def chart_preview(
    sid: str,
    viz_type: str = Query(...),
    metric_col: str = Query(...),
    aggregate: str = Query("SUM"),
    dimension_col: str | None = Query(default=None),
    time_col: str | None = Query(default=None),
    time_grain: str = Query("P1M"),
    row_limit: int = Query(10),
):
    sess = get_session(sid)
    dataset_info = sess["phase3"].get("dataset_info")
    if dataset_info is None:
        raise HTTPException(status_code=400, detail="No dataset loaded yet. Run Phase 3 plan first.")

    from tools.superset_api import SupersetClient
    from tools.column_sampler import ColumnSampler

    sup_cfg = sess["superset"]
    try:
        client = SupersetClient(
            base_url=sup_cfg["url"] or config.SUPERSET_URL,
            username=sup_cfg["username"] or config.SUPERSET_USERNAME,
            password=sup_cfg["password"] or config.SUPERSET_PASSWORD,
            session_cookie=sup_cfg.get("session_cookie") or None,
            csrf_token=sup_cfg.get("csrf_token") or None,
        )
        client.authenticate()
    except Exception as e:
        return {"rows": [], "total": None, "error": f"Superset auth failed: {e}"}

    sampler = ColumnSampler(client)
    table_ref = sampler._get_table_ref(dataset_info)
    database_id = sampler._get_database_id(dataset_info.id)
    if database_id is None:
        return {"rows": [], "total": None, "error": "Could not determine database ID"}

    def _q(name: str) -> str:
        return f'"{name.replace(chr(34), chr(34) + chr(34))}"'

    aggregate = aggregate.upper()
    allowed_aggregates = {"SUM", "AVG", "COUNT", "MIN", "MAX"}
    if aggregate not in allowed_aggregates:
        return {"rows": [], "total": None, "error": f"Unsupported aggregate '{aggregate}'"}

    grain_map = {"P1D": "day", "P1W": "week", "P1M": "month", "P1Y": "year"}
    grain_sql = grain_map.get(time_grain, "month")

    try:
        if viz_type == "big_number_total" or (not dimension_col and not time_col):
            sql = f"SELECT {aggregate}({_q(metric_col)}) AS value FROM {table_ref}"
            resp = client._request(
                "POST",
                "/api/v1/sqllab/execute/",
                json={"database_id": database_id, "sql": sql, "runAsync": False, "queryLimit": 1},
            )
            rows_data = resp.json().get("data", [])
            total_val = rows_data[0].get("value") if rows_data else None
            total = float(total_val) if total_val is not None else 0.0
            return {"rows": [], "total": total, "error": None}
        elif viz_type == "echarts_timeseries_line" and time_col:
            sql = (
                f"SELECT DATE_TRUNC('{grain_sql}', {_q(time_col)}) AS label, "
                f"{aggregate}({_q(metric_col)}) AS value "
                f"FROM {table_ref} "
                f"GROUP BY 1 ORDER BY 1 LIMIT {row_limit}"
            )
        else:
            dim = dimension_col or metric_col
            sql = (
                f"SELECT {_q(dim)} AS label, "
                f"{aggregate}({_q(metric_col)}) AS value "
                f"FROM {table_ref} "
                f"GROUP BY {_q(dim)} "
                f"ORDER BY 2 DESC LIMIT {row_limit}"
            )

        resp = client._request(
            "POST",
            "/api/v1/sqllab/execute/",
            json={"database_id": database_id, "sql": sql, "runAsync": False, "queryLimit": row_limit},
        )
        rows_data = resp.json().get("data", [])
        rows = []
        for r in rows_data:
            vals = list(r.values())
            label = str(vals[0]) if vals[0] is not None else "null"
            value = float(vals[1]) if len(vals) > 1 and vals[1] is not None else 0.0
            rows.append({"label": label, "value": value})
        return {"rows": rows, "total": None, "error": None}
    except Exception as e:
        return {"rows": [], "total": None, "error": str(e)}


@app.post("/api/sessions/{sid}/phase3/plan/update")
async def update_plan(sid: str, body: UpdatePlanBody):
    sess = get_session(sid)
    from models.schemas import DashboardPlan
    plan_dict = dict(body.plan) if body.plan else {}
    # Ensure chart specs include required `filters` field (absent in TypeScript interface)
    for chart in plan_dict.get("charts", []):
        if isinstance(chart, dict):
            chart.setdefault("filters", [])
            chart.setdefault("row_limit", None)
            chart.setdefault("sort_by", None)
    plan_dict.setdefault("position_json", {})
    sess["phase3"]["dashboard_plan"] = DashboardPlan(**plan_dict)
    return {"ok": True}


def _metric_description_label(metric: Any) -> str:
    if isinstance(metric, dict):
        label = metric.get("label")
        if label:
            return str(label)
        column = metric.get("column")
        if isinstance(column, dict):
            return str(column.get("column_name") or column.get("name") or "")
        if column:
            return str(column)
        aggregate = metric.get("aggregate")
        if aggregate:
            return str(aggregate)
    return str(metric)


def _fallback_chart_description(spec: ChartDescriptionSpec, dataset_name: str) -> str:
    title = spec.title.lower().strip() or "this chart"
    dataset = dataset_name.strip() or "the dataset"
    return f"Shows {title} for the {dataset} dataset."


def _build_chart_description_prompt(
    specs: list[ChartDescriptionSpec],
    dataset_name: str,
    dashboard_title: str,
) -> tuple[str, str]:
    grain_map = {"P1D": "daily", "P1W": "weekly", "P1M": "monthly", "P1Y": "yearly"}
    chart_lines = []
    for idx, spec in enumerate(specs, 1):
        metrics_str = ", ".join(
            label for label in (_metric_description_label(m) for m in spec.metrics) if label
        ) or "none"
        groupby_str = ", ".join(spec.groupby) if spec.groupby else "none"
        time_str = ""
        if spec.time_column:
            grain = spec.time_grain or ""
            grain_label = grain_map.get(grain, grain)
            time_str = f", time: {spec.time_column} ({grain_label})"
        chart_lines.append(
            f"{idx}. Title: {spec.title}\n"
            f"   Type: {spec.viz_type}\n"
            f"   Metrics: {metrics_str}\n"
            f"   Group by: {groupby_str}{time_str}\n"
        )

    system = (
        "You are a data analyst writing chart descriptions for a business dashboard. "
        "Write exactly one sentence per chart that explains what the chart shows and why it is useful. "
        "Be specific - mention the metric name and dimension. Use plain business language, no technical jargon. "
        "Keep each description under 20 words.\n\n"
        "Return ONLY valid JSON - no markdown, no preamble:\n"
        "{\n"
        '  "descriptions": [\n'
        '    "description for chart 1",\n'
        '    "description for chart 2"\n'
        "  ]\n"
        "}\n\n"
        f"The array must have exactly {len(specs)} items in the same order as the input list."
    )
    user = (
        f"Dashboard: {dashboard_title}\n"
        f"Dataset: {dataset_name}\n\n"
        "Charts:\n"
        f"{''.join(chart_lines)}\n"
        "Write one description per chart."
    )
    return system, user


def _parse_description_response(raw: str) -> list[str]:
    data = json.loads(raw)
    descriptions = data.get("descriptions")
    if not isinstance(descriptions, list):
        raise ValueError("LLM response missing descriptions array")
    return [str(item).strip() for item in descriptions]


def _generate_chart_descriptions(
    specs: list[ChartDescriptionSpec],
    dataset_name: str,
    dashboard_title: str,
) -> list[str]:
    if not specs:
        return []
    from tools.llm_client import chat

    system, user = _build_chart_description_prompt(specs, dataset_name, dashboard_title)
    descriptions: list[str] = []
    for attempt in range(2):
        try:
            raw = chat(system=system, user=user, max_tokens=1600, temperature=0)
            descriptions = _parse_description_response(raw)
            break
        except Exception:
            if attempt == 1:
                descriptions = [
                    _fallback_chart_description(spec, dataset_name)
                    for spec in specs
                ]

    while len(descriptions) < len(specs):
        descriptions.append(_fallback_chart_description(specs[len(descriptions)], dataset_name))
    return descriptions[:len(specs)]


@app.post("/api/sessions/{sid}/phase3/generate-descriptions")
async def generate_chart_descriptions(sid: str, body: GenerateDescriptionsBody):
    try:
        sess = get_session(sid)
        if sess.get("llm_model"):
            config.LLM_MODEL = sess["llm_model"]

        client = _superset_client_from_session(sess)
        descriptions = _generate_chart_descriptions(
            body.chart_specs,
            body.dataset_name,
            body.dashboard_title,
        )

        results = []
        for idx, spec in enumerate(body.chart_specs):
            chart_id = body.chart_ids[idx] if idx < len(body.chart_ids) else spec.id
            description = descriptions[idx] if idx < len(descriptions) else _fallback_chart_description(
                spec,
                body.dataset_name,
            )
            if not chart_id:
                results.append(
                    {
                        "id": chart_id,
                        "title": spec.title,
                        "description": description,
                        "ok": False,
                        "error": "Missing chart id",
                    }
                )
                continue
            try:
                client._request(
                    "PUT",
                    f"/api/v1/chart/{chart_id}",
                    json={"description": description},
                )
                results.append(
                    {
                        "id": chart_id,
                        "title": spec.title,
                        "description": description,
                        "ok": True,
                        "error": None,
                    }
                )
            except Exception as e:
                results.append(
                    {
                        "id": chart_id,
                        "title": spec.title,
                        "description": description,
                        "ok": False,
                        "error": str(e),
                    }
                )

        succeeded = sum(1 for result in results if result["ok"])
        return {
            "results": results,
            "total": len(results),
            "succeeded": succeeded,
            "failed": len(results) - succeeded,
            "error": None,
        }
    except Exception as e:
        total = len(body.chart_specs)
        return {
            "results": [],
            "total": total,
            "succeeded": 0,
            "failed": total,
            "error": str(e),
        }


# ── Data Quality endpoints ────────────────────────────────────────────────────

@app.get("/api/sessions/{sid}/phase3/data-quality")
async def get_data_quality_report(sid: str, dataset_name: str = Query(...)):
    sess = get_session(sid)

    def worker():
        import pandas as pd
        import numpy as np
        from tools.superset_api import SupersetClient
        from tools.llm_client import chat

        sup_cfg = sess["superset"]
        client = SupersetClient(
            base_url=sup_cfg["url"] or config.SUPERSET_URL,
            username=sup_cfg["username"] or config.SUPERSET_USERNAME,
            password=sup_cfg["password"] or config.SUPERSET_PASSWORD,
            session_cookie=sup_cfg.get("session_cookie") or None,
            csrf_token=sup_cfg.get("csrf_token") or None,
        )
        client.authenticate()

        dataset = client.get_dataset_by_name(dataset_name)
        if not dataset:
            raise ValueError(f"Dataset '{dataset_name}' not found in Superset")

        dataset_id = dataset.id
        db_id = client.get_dataset_database_id(dataset_id)
        if not db_id:
            raise ValueError("Could not determine database ID for dataset")

        from tools.column_sampler import ColumnSampler
        sampler = ColumnSampler(client)
        table_ref = sampler._get_table_ref(dataset)
        sql = f"SELECT * FROM {table_ref} LIMIT 1000"

        ok, rows, err = client.execute_sql(db_id, sql, limit=1000)
        if not ok or not rows:
            raise ValueError(f"Could not fetch sample data: {err}")

        df = pd.DataFrame(rows)
        total_rows = len(df)
        total_cols = len(df.columns)

        # Duplicate detection
        dup_count = int(df.duplicated().sum())
        dup_pct = round(dup_count / total_rows * 100, 2) if total_rows > 0 else 0.0

        # Column-level null stats
        cols_with_nulls = int((df.isnull().sum() > 0).sum())
        cols_all_null = int((df.isnull().sum() == total_rows).sum())
        cols_single_value = int((df.nunique() <= 1).sum())

        dataset_metrics = {
            "total_rows": total_rows,
            "total_columns": total_cols,
            "duplicate_rows_count": dup_count,
            "duplicate_rows_pct": dup_pct,
            "columns_with_nulls": cols_with_nulls,
            "columns_all_null": cols_all_null,
            "columns_single_value": cols_single_value,
            "sample_size": total_rows,
        }

        issues = []
        column_profiles = []

        if dup_count > 0:
            issues.append({
                "severity": "critical" if dup_pct > 10 else "warning",
                "column": None,
                "issue_type": "duplicates",
                "description": f"{dup_count} duplicate rows ({dup_pct}% of sample)",
                "affected_rows": dup_count,
                "affected_pct": dup_pct,
            })

        for col in df.columns:
            series = df[col]
            null_count = int(series.isnull().sum())
            null_pct = round(null_count / total_rows * 100, 2) if total_rows > 0 else 0.0
            distinct_count = int(series.nunique(dropna=True))
            distinct_pct = round(distinct_count / total_rows * 100, 2) if total_rows > 0 else 0.0

            profile: dict = {
                "name": col,
                "dtype": str(series.dtype),
                "null_count": null_count,
                "null_pct": null_pct,
                "distinct_count": distinct_count,
                "distinct_pct": distinct_pct,
                "min_val": None,
                "max_val": None,
                "mean_val": None,
                "std_dev": None,
                "outlier_count": None,
                "outlier_pct": None,
                "zeros_pct": None,
                "top_values": None,
                "avg_length": None,
                "min_date": None,
                "max_date": None,
                "date_range_days": None,
                "future_dates_count": None,
            }

            if null_pct > 50:
                issues.append({
                    "severity": "critical",
                    "column": col,
                    "issue_type": "high_nulls",
                    "description": f"Column '{col}' has {null_pct}% null values",
                    "affected_rows": null_count,
                    "affected_pct": null_pct,
                })
            elif null_pct > 20:
                issues.append({
                    "severity": "warning",
                    "column": col,
                    "issue_type": "moderate_nulls",
                    "description": f"Column '{col}' has {null_pct}% null values",
                    "affected_rows": null_count,
                    "affected_pct": null_pct,
                })
            elif null_pct > 5:
                issues.append({
                    "severity": "info",
                    "column": col,
                    "issue_type": "low_nulls",
                    "description": f"Column '{col}' has {null_pct}% null values",
                    "affected_rows": null_count,
                    "affected_pct": null_pct,
                })

            if distinct_count <= 1 and total_rows > 1:
                issues.append({
                    "severity": "warning",
                    "column": col,
                    "issue_type": "single_value",
                    "description": f"Column '{col}' has only {distinct_count} distinct value(s) — may be constant",
                    "affected_rows": total_rows,
                    "affected_pct": 100.0,
                })

            numeric_series = pd.to_numeric(series, errors="coerce")
            is_numeric = numeric_series.notna().sum() > total_rows * 0.8

            is_date = False
            if not is_numeric:
                try:
                    date_series = pd.to_datetime(series.dropna(), errors="coerce", infer_datetime_format=True)
                    if date_series.notna().sum() > len(series.dropna()) * 0.8:
                        is_date = True
                except Exception:
                    pass

            if is_numeric:
                clean = numeric_series.dropna()
                if len(clean) > 0:
                    mn = float(clean.min())
                    mx = float(clean.max())
                    mean = float(clean.mean())
                    std = float(clean.std()) if len(clean) > 1 else 0.0
                    profile["min_val"] = mn
                    profile["max_val"] = mx
                    profile["mean_val"] = round(mean, 4)
                    profile["std_dev"] = round(std, 4)

                    zeros = int((clean == 0).sum())
                    profile["zeros_pct"] = round(zeros / len(clean) * 100, 2)

                    if std > 0:
                        z_scores = np.abs((clean - mean) / std)
                        outliers = int((z_scores > 3).sum())
                        outlier_pct = round(outliers / len(clean) * 100, 2)
                        profile["outlier_count"] = outliers
                        profile["outlier_pct"] = outlier_pct
                        if outliers > 0 and outlier_pct > 5:
                            issues.append({
                                "severity": "warning",
                                "column": col,
                                "issue_type": "outliers",
                                "description": f"Column '{col}' has {outliers} outliers ({outlier_pct}%) beyond 3σ",
                                "affected_rows": outliers,
                                "affected_pct": outlier_pct,
                            })

            elif is_date:
                try:
                    date_series = pd.to_datetime(series.dropna(), errors="coerce", infer_datetime_format=True)
                    clean_dates = date_series.dropna()
                    if len(clean_dates) > 0:
                        profile["min_date"] = str(clean_dates.min().date())
                        profile["max_date"] = str(clean_dates.max().date())
                        profile["date_range_days"] = int((clean_dates.max() - clean_dates.min()).days)
                        future = int((clean_dates > pd.Timestamp.now()).sum())
                        profile["future_dates_count"] = future
                        if future > 0:
                            issues.append({
                                "severity": "warning",
                                "column": col,
                                "issue_type": "future_dates",
                                "description": f"Column '{col}' has {future} future date values",
                                "affected_rows": future,
                                "affected_pct": round(future / len(clean_dates) * 100, 2),
                            })
                except Exception:
                    pass

            else:
                str_series = series.dropna().astype(str)
                if len(str_series) > 0:
                    avg_len = round(float(str_series.str.len().mean()), 2)
                    profile["avg_length"] = avg_len

                value_counts = series.value_counts().head(5)
                top_vals = [
                    {"value": str(v), "count": int(c), "pct": round(int(c) / total_rows * 100, 2)}
                    for v, c in value_counts.items()
                ]
                profile["top_values"] = top_vals

            column_profiles.append(profile)

        critical_count = sum(1 for i in issues if i["severity"] == "critical")
        warning_count = sum(1 for i in issues if i["severity"] == "warning")
        info_count = sum(1 for i in issues if i["severity"] == "info")
        score = max(0, min(100, 100 - 20 * critical_count - 5 * warning_count - 1 * info_count))

        if score >= 90:
            grade = "Excellent"
        elif score >= 75:
            grade = "Good"
        elif score >= 50:
            grade = "Fair"
        elif score >= 25:
            grade = "Poor"
        else:
            grade = "Critical"

        llm_summary = ""
        try:
            import concurrent.futures
            issue_lines = "\n".join(
                f"- [{i['severity'].upper()}] {i['description']}" for i in issues[:10]
            )
            null_lines = "\n".join(
                f"- {p['name']}: {p['null_pct']}% null"
                for p in column_profiles if p["null_pct"] > 0
            )[:500]
            prompt = (
                f"Dataset: {dataset_name}\n"
                f"Rows sampled: {total_rows}, Columns: {total_cols}\n"
                f"Duplicates: {dup_count} ({dup_pct}%)\n"
                f"Quality score: {score}/100 ({grade})\n\n"
                f"Issues found:\n{issue_lines or 'None'}\n\n"
                f"Null summary:\n{null_lines or 'None'}\n\n"
                "Write a concise 2-3 sentence plain-English data quality summary for a business analyst. "
                "Focus on the most important issues and their potential impact on dashboard accuracy."
            )
            system = "You are a data quality analyst. Write clear, actionable summaries."
            with concurrent.futures.ThreadPoolExecutor(max_workers=1) as ex:
                future = ex.submit(chat, system, prompt, 200, 0)
                llm_summary = future.result(timeout=10)
        except Exception:
            llm_summary = ""

        report = {
            "dataset_name": dataset_name,
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "overall_score": score,
            "overall_grade": grade,
            "summary": llm_summary,
            "dataset_metrics": dataset_metrics,
            "issues": issues,
            "column_profiles": column_profiles,
            "error": None,
        }

        sess["phase3"]["dq_report"] = report
        append_audit(sess, 3, "dq_report_generated", "Data Quality Report",
                     f"Score: {score}/100 ({grade}), {len(issues)} issues found", status="info")
        return report

    loop = asyncio.get_event_loop()
    try:
        report = await loop.run_in_executor(None, worker)
        return report
    except Exception as e:
        return {
            "dataset_name": dataset_name,
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "overall_score": 0,
            "overall_grade": "Critical",
            "summary": "",
            "dataset_metrics": {
                "total_rows": 0, "total_columns": 0, "duplicate_rows_count": 0,
                "duplicate_rows_pct": 0.0, "columns_with_nulls": 0, "columns_all_null": 0,
                "columns_single_value": 0, "sample_size": 0,
            },
            "issues": [],
            "column_profiles": [],
            "error": str(e),
        }


@app.post("/api/sessions/{sid}/phase3/data-quality/export-pdf")
async def export_dq_report_pdf(sid: str, report: dict):
    get_session(sid)

    def build_pdf(r: dict) -> bytes:
        try:
            from weasyprint import HTML
        except ImportError:
            raise RuntimeError("weasyprint is not installed. Run: pip install weasyprint")

        score = r.get("overall_score", 0)
        grade = r.get("overall_grade", "N/A")
        dataset_name = r.get("dataset_name", "")
        generated_at = r.get("generated_at", "")
        summary = r.get("summary", "")
        metrics = r.get("dataset_metrics", {})
        issues = r.get("issues", [])
        profiles = r.get("column_profiles", [])

        if score >= 90:
            score_color = "#22c55e"
        elif score >= 75:
            score_color = "#14b8a6"
        elif score >= 50:
            score_color = "#f59e0b"
        elif score >= 25:
            score_color = "#f97316"
        else:
            score_color = "#ef4444"

        sev_colors = {"critical": "#ef4444", "warning": "#f59e0b", "info": "#6366f1"}

        issues_html = ""
        for sev in ("critical", "warning", "info"):
            grp = [i for i in issues if i.get("severity") == sev]
            if not grp:
                continue
            issues_html += f'<h3 style="color:{sev_colors[sev]};text-transform:capitalize">{sev}</h3><ul>'
            for iss in grp:
                issues_html += f'<li>{iss.get("description","")}</li>'
            issues_html += "</ul>"

        profiles_rows = ""
        for p in profiles:
            null_pct = p.get("null_pct", 0)
            null_style = ""
            if null_pct > 20:
                null_style = "color:#ef4444;font-weight:bold"
            elif null_pct > 5:
                null_style = "color:#f59e0b;font-weight:bold"
            profiles_rows += (
                f'<tr><td>{p.get("name","")}</td><td>{p.get("dtype","")}</td>'
                f'<td style="{null_style}">{null_pct}%</td>'
                f'<td>{p.get("distinct_count","")}</td>'
                f'<td>{p.get("outlier_count") if p.get("outlier_count") is not None else "-"}</td></tr>'
            )

        html = f"""<!DOCTYPE html><html><head>
<meta charset="utf-8">
<style>
  body{{font-family:Arial,sans-serif;font-size:12px;color:#111;margin:20px}}
  h1{{font-size:20px}} h2{{font-size:15px;border-bottom:1px solid #ccc;padding-bottom:4px}}
  .score-circle{{display:inline-block;width:60px;height:60px;border-radius:50%;
    background:{score_color};color:#fff;text-align:center;line-height:60px;
    font-size:20px;font-weight:bold;margin-right:12px;vertical-align:middle}}
  .metric-grid{{display:grid;grid-template-columns:repeat(4,1fr);gap:10px;margin:10px 0}}
  .metric-card{{background:#f5f5f5;border-radius:6px;padding:10px;text-align:center}}
  .metric-card .val{{font-size:20px;font-weight:bold}}
  table{{width:100%;border-collapse:collapse;margin-top:8px}}
  th,td{{border:1px solid #ddd;padding:5px 8px;text-align:left}}
  th{{background:#f0f0f0}}
  .summary-box{{background:#f9f9f9;border-left:4px solid {score_color};padding:10px;margin:10px 0}}
  footer{{color:#888;font-size:10px;margin-top:20px;border-top:1px solid #ddd;padding-top:8px}}
</style></head><body>
<h1><span class="score-circle">{score}</span> Data Quality Report — {dataset_name}</h1>
<p><strong>Grade:</strong> {grade} &nbsp;|&nbsp; <strong>Generated:</strong> {generated_at}</p>
{f'<div class="summary-box">{summary}</div>' if summary else ''}
<h2>Dataset Overview</h2>
<div class="metric-grid">
  <div class="metric-card"><div class="val">{metrics.get("total_rows",0):,}</div>Rows</div>
  <div class="metric-card"><div class="val">{metrics.get("total_columns",0)}</div>Columns</div>
  <div class="metric-card"><div class="val">{metrics.get("duplicate_rows_pct",0)}%</div>Duplicates</div>
  <div class="metric-card"><div class="val">{metrics.get("columns_with_nulls",0)}</div>Cols w/ Nulls</div>
</div>
<h2>Issues ({len(issues)} total)</h2>
{issues_html or "<p>No issues found.</p>"}
<h2>Column Profiles</h2>
<table><thead><tr><th>Column</th><th>Type</th><th>Null %</th><th>Distinct</th><th>Outliers</th></tr></thead>
<tbody>{profiles_rows}</tbody></table>
<footer>Generated by DATA VIZ — Superset Dashboard Builder</footer>
</body></html>"""

        return HTML(string=html).write_pdf()

    loop = asyncio.get_event_loop()
    pdf_bytes = await loop.run_in_executor(None, build_pdf, report)
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": "attachment; filename=data-quality-report.pdf"},
    )


# ── Dashboard version history endpoints ───────────────────────────────────────

@app.get("/api/sessions/{sid}/dashboard-history")
async def dashboard_history_list(sid: str):
    get_session(sid)
    from tools.version_history import DashboardVersionHistory

    return {"dashboards": DashboardVersionHistory().list_all_dashboards()}


@app.get("/api/sessions/{sid}/dashboard-history/{dashboard_id}")
async def dashboard_history_index(sid: str, dashboard_id: int):
    get_session(sid)
    from tools.version_history import DashboardVersionHistory

    index = DashboardVersionHistory().load_index(dashboard_id)
    if not index:
        raise HTTPException(status_code=404, detail="No history found")
    return index


@app.get("/api/sessions/{sid}/dashboard-history/{dashboard_id}/versions/{version}")
async def dashboard_history_snapshot(sid: str, dashboard_id: int, version: int):
    get_session(sid)
    from tools.version_history import DashboardVersionHistory

    snapshot = DashboardVersionHistory().load_snapshot(dashboard_id, version)
    if not snapshot:
        raise HTTPException(status_code=404, detail="Version not found")
    return snapshot


@app.post("/api/sessions/{sid}/dashboard-history/{dashboard_id}/restore/{version}")
async def dashboard_history_restore(sid: str, dashboard_id: int, version: int):
    sess = get_session(sid)
    from tools.version_history import DashboardVersionHistory

    vh = DashboardVersionHistory()
    snapshot = vh.load_snapshot(dashboard_id, version)
    if not snapshot:
        raise HTTPException(status_code=404, detail="Version not found")

    try:
        client = _superset_client_from_session(sess)
    except Exception as e:
        return {"ok": False, "error": f"Superset auth failed: {e}"}

    backup_version = -1
    try:
        current_state = client.get_dashboard_full_state(dashboard_id)
        if current_state:
            backup_version = _save_full_state_snapshot(
                version_history=vh,
                dashboard_id=dashboard_id,
                state=current_state,
                dashboard_url=f"{client.base_url}/superset/dashboard/{dashboard_id}",
                requirements_prompt=snapshot.get("requirements_prompt", ""),
                dataset_name=snapshot.get("dataset_name", ""),
                created_by=_created_by_from_session(sess),
                change_summary=f"Before restore to version {version}",
            )
    except Exception:
        backup_version = -1

    success = client.restore_dashboard_version(dashboard_id, snapshot)
    if not success:
        return {"ok": False, "error": "Restore failed"}

    saved_version = vh.save_snapshot(
        dashboard_id=dashboard_id,
        dashboard_title=snapshot.get("dashboard_title", ""),
        dashboard_url=snapshot.get("dashboard_url", ""),
        requirements_prompt=snapshot.get("requirements_prompt", ""),
        dataset_name=snapshot.get("dataset_name", ""),
        position_json=snapshot.get("position_json", {}),
        charts=snapshot.get("charts", []),
        filters=snapshot.get("filters", []),
        created_by=_created_by_from_session(sess),
        change_summary=f"Restored from version {version}",
    )
    return {
        "ok": True,
        "message": f"Restored to v{version}",
        "saved_version": saved_version,
        "backup_version": backup_version,
    }


@app.delete("/api/sessions/{sid}/dashboard-history/{dashboard_id}/versions/{version}")
async def dashboard_history_delete(sid: str, dashboard_id: int, version: int):
    get_session(sid)
    from tools.version_history import DashboardVersionHistory

    success = DashboardVersionHistory().delete_version(dashboard_id, version)
    return {"ok": success}


@app.get("/api/sessions/{sid}/dashboard-history/{dashboard_id}/versions/{version}/download")
async def dashboard_history_download(sid: str, dashboard_id: int, version: int):
    get_session(sid)
    from tools.version_history import DashboardVersionHistory

    snapshot = DashboardVersionHistory().load_snapshot(dashboard_id, version)
    if not snapshot:
        raise HTTPException(status_code=404, detail="Version not found")
    return JSONResponse(
        content=snapshot,
        headers={
            "Content-Disposition": (
                f'attachment; filename="dashboard_{dashboard_id}_v{version}.json"'
            )
        },
    )


# ── CSV analysis endpoints ────────────────────────────────────────────────────

_CSV_MAX_BYTES = 50 * 1024 * 1024  # 50 MB
_CSV_CHART_ROW_LIMIT = 100


@app.post("/api/sessions/{sid}/csv/upload")
async def csv_upload(sid: str, file: UploadFile = File(...)):
    sess = get_session(sid)
    filename = file.filename or "upload"
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext not in ("csv", "xlsx", "xls"):
        raise HTTPException(status_code=400, detail="Only .csv / .xlsx / .xls files are supported")
    content = await file.read()
    if len(content) > _CSV_MAX_BYTES:
        raise HTTPException(status_code=400, detail="File exceeds 50 MB limit")
    try:
        import io
        import pandas as pd

        if ext == "csv":
            df = pd.read_csv(io.BytesIO(content))
        else:
            df = pd.read_excel(io.BytesIO(content))

        columns = []
        for col in df.columns:
            dtype = str(df[col].dtype)
            sample = [str(v) for v in df[col].dropna().head(5).tolist()]
            columns.append({"name": col, "dtype": dtype, "sample": sample})

        sess["csv_session"] = {
            "filename": filename,
            "row_count": len(df),
            "columns": columns,
            "df_json": df.to_json(orient="records"),
            "charts": [],
        }
        return {
            "filename": filename,
            "row_count": len(df),
            "columns": columns,
            "charts": [],
        }
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"Failed to parse file: {e}")


@app.post("/api/sessions/{sid}/csv/query")
async def csv_query(sid: str, body: CsvQueryBody):
    sess = get_session(sid)
    csv_sess = sess.get("csv_session")
    if not csv_sess:
        raise HTTPException(status_code=400, detail="No CSV loaded. Upload a file first.")
    if sess.get("llm_model"):
        config.LLM_MODEL = sess["llm_model"]
    try:
        import duckdb
        import pandas as pd
        import io
        from tools.llm_client import chat

        df = pd.read_json(io.StringIO(csv_sess["df_json"]))
        columns_desc = ", ".join(
            f"{c['name']} ({c['dtype']})" for c in csv_sess["columns"]
        )
        history_text = ""
        for turn in body.history[-6:]:
            role = turn.get("role", "user")
            content = turn.get("content", "")
            history_text += f"\n{role.upper()}: {content}"

        system = (
            "You are a data analyst. Given a DuckDB table named `df` with these columns:\n"
            f"{columns_desc}\n\n"
            "Return a JSON object with exactly these keys:\n"
            '{"sql": "<SELECT query on df>", "chart_type": "<bar|line|pie|table|big_number>"}\n'
            "Rules:\n"
            "- Always query the table named `df`\n"
            "- chart_type=big_number when the result is a single aggregate value\n"
            "- chart_type=pie for proportions with ≤10 categories\n"
            "- chart_type=line for time-series data\n"
            "- chart_type=bar for categorical comparisons\n"
            "- chart_type=table when the user wants raw rows\n"
            "- Limit results to 100 rows unless user asks for more\n"
            "- Return ONLY the JSON, no markdown fences"
        )
        user_prompt = f"Previous conversation:{history_text}\n\nNew question: {body.question}"

        raw = chat(system=system, user=user_prompt, temperature=0)
        raw = raw.strip()
        if raw.startswith("```"):
            raw = raw.split("```")[1]
            if raw.startswith("json"):
                raw = raw[4:]
        parsed = json.loads(raw)
        sql = parsed["sql"].strip()
        chart_type = parsed.get("chart_type", "table")
        if not _is_readonly_select_sql(sql):
            raise ValueError("Only read-only SELECT queries are allowed")

        con = duckdb.connect()
        con.register("df", df)
        result = con.execute(sql).fetchdf()
        con.close()

        rows = result.head(_CSV_CHART_ROW_LIMIT).to_dict(orient="records")
        return {
            "question": body.question,
            "sql": sql,
            "columns": list(result.columns),
            "rows": rows,
            "row_count": len(result),
            "chart_type": chart_type,
            "error": None,
        }
    except Exception as e:
        return {
            "question": body.question,
            "sql": "",
            "columns": [],
            "rows": [],
            "row_count": 0,
            "chart_type": "table",
            "error": str(e),
        }


@app.post("/api/sessions/{sid}/csv/add-chart")
async def csv_add_chart(sid: str, body: CsvChartBody):
    sess = get_session(sid)
    csv_sess = sess.get("csv_session")
    if not csv_sess:
        raise HTTPException(status_code=400, detail="No CSV session")
    chart = body.model_dump()
    # Deduplicate by id
    csv_sess["charts"] = [c for c in csv_sess["charts"] if c["id"] != chart["id"]]
    csv_sess["charts"].append(chart)
    return {"ok": True}


@app.delete("/api/sessions/{sid}/csv/charts/{chart_id}")
async def csv_remove_chart(sid: str, chart_id: str):
    sess = get_session(sid)
    csv_sess = sess.get("csv_session")
    if not csv_sess:
        raise HTTPException(status_code=400, detail="No CSV session")
    csv_sess["charts"] = [c for c in csv_sess["charts"] if c["id"] != chart_id]
    return {"ok": True}


@app.get("/api/sessions/{sid}/csv/charts")
async def csv_get_charts(sid: str):
    sess = get_session(sid)
    csv_sess = sess.get("csv_session")
    if not csv_sess:
        return {"charts": []}
    return {"charts": csv_sess.get("charts", [])}


@app.post("/api/sessions/{sid}/csv/export/pdf")
async def csv_export_pdf(sid: str):
    sess = get_session(sid)
    csv_sess = sess.get("csv_session")
    if not csv_sess:
        raise HTTPException(status_code=400, detail="No CSV session")
    try:
        from weasyprint import HTML

        charts = csv_sess.get("charts", [])
        filename = csv_sess.get("filename", "data")

        def _table_html(chart: dict) -> str:
            cols = chart.get("columns", [])
            rows = chart.get("rows", [])
            header = "".join(f"<th>{c}</th>" for c in cols)
            body_rows = ""
            for row in rows[:50]:
                cells = "".join(f"<td>{row.get(c, '')}</td>" for c in cols)
                body_rows += f"<tr>{cells}</tr>"
            return f"<table><thead><tr>{header}</tr></thead><tbody>{body_rows}</tbody></table>"

        cards = ""
        for chart in charts:
            cards += (
                f'<div class="chart-card">'
                f'<h3>{chart.get("title", "Chart")}</h3>'
                f'<p class="meta">Type: {chart.get("chart_type")} · '
                f'Question: {chart.get("question", "")}</p>'
                f'{_table_html(chart)}'
                f'</div>'
            )

        html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8">
<style>
body {{ font-family: sans-serif; margin: 20px; color: #111; }}
h1 {{ font-size: 1.4em; border-bottom: 2px solid #6366f1; padding-bottom: 8px; }}
.chart-card {{ margin-bottom: 24px; page-break-inside: avoid; }}
h3 {{ font-size: 1em; margin: 0 0 4px; }}
.meta {{ font-size: 0.75em; color: #555; margin: 0 0 8px; }}
table {{ border-collapse: collapse; width: 100%; font-size: 0.8em; }}
th {{ background: #f0f0f0; border: 1px solid #ddd; padding: 4px 8px; text-align: left; }}
td {{ border: 1px solid #ddd; padding: 4px 8px; }}
tr:nth-child(even) {{ background: #f9f9f9; }}
</style></head>
<body>
<h1>Dashboard Export — {filename}</h1>
{cards}
</body></html>"""

        pdf_bytes = HTML(string=html).write_pdf()
        return Response(
            content=pdf_bytes,
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="dashboard.pdf"'},
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"PDF generation failed: {e}")


@app.post("/api/sessions/{sid}/csv/export/excel")
async def csv_export_excel(sid: str):
    sess = get_session(sid)
    csv_sess = sess.get("csv_session")
    if not csv_sess:
        raise HTTPException(status_code=400, detail="No CSV session")
    try:
        import io as _io
        import openpyxl
        from openpyxl.styles import Font, PatternFill

        charts = csv_sess.get("charts", [])
        wb = openpyxl.Workbook()
        wb.remove(wb.active)  # remove default sheet

        def _safe_sheet_name(name: str) -> str:
            for ch in r'\/?*[]:"':
                name = name.replace(ch, "_")
            return name[:31]

        for chart in charts:
            title = chart.get("title", "Chart")
            ws = wb.create_sheet(title=_safe_sheet_name(title))
            cols = chart.get("columns", [])
            rows = chart.get("rows", [])
            # Header row
            for ci, col in enumerate(cols, start=1):
                cell = ws.cell(row=1, column=ci, value=col)
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="6366F1")
                cell.font = Font(bold=True, color="FFFFFF")
            for ri, row in enumerate(rows, start=2):
                for ci, col in enumerate(cols, start=1):
                    ws.cell(row=ri, column=ci, value=row.get(col, ""))

        buf = _io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return Response(
            content=buf.read(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="dashboard.xlsx"'},
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Excel generation failed: {e}")


# ── Audit log endpoints ───────────────────────────────────────────────────────

@app.get("/api/sessions/{sid}/audit")
async def get_audit_log(sid: str):
    sess = get_session(sid)
    return {"entries": sess.get("audit_log", [])}


@app.delete("/api/sessions/{sid}/audit")
async def clear_audit_log(sid: str):
    sess = get_session(sid)
    sess["audit_log"] = []
    return {"ok": True}


if __name__ == "__main__":
    import uvicorn
    uvicorn.run(app, host="0.0.0.0", port=8000)
